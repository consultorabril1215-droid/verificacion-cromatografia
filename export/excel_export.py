"""
Exportación a Excel del resultado de verificación, con **fórmulas vivas**
(no solo valores calculados): un auditor que abra el archivo puede dar clic
en cualquier celda de resultado y ver la fórmula que la sustenta, referida a
los datos crudos que también quedan en la misma hoja — igual que en los
formatos LA-F-319/LA-F-210 originales del laboratorio.

Excepción documentada: la prueba de linealidad de Mandel (ISO 8466-1 §4.1.3,
que requiere un ajuste cuadrático) y la regresión ponderada (cuando la curva
no es homogénea) se calculan en Python y se insertan como valores, porque
Excel no tiene una función nativa simple de regresión polinómica/ponderada
sin fórmulas matriciales; se deja una nota junto al resultado indicándolo.
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from core.models import VerificationProject
from core.grubbs_table import GRUBBS_CRITICAL
from core.dixon_table import DIXON_CRITICAL_95
from core import stats as S
from core import uncertainty as U

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
FORMULA_FONT = Font(color="0000CC")   # celdas con fórmula, en azul, para distinguirlas de datos crudos


def _hdr(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    return c


def _formula(ws, row, col, formula, fmt=None):
    c = ws.cell(row=row, column=col, value=formula)
    c.font = FORMULA_FONT
    if fmt:
        c.number_format = fmt
    return c


def _autosize(ws, n_cols):
    for i in range(1, n_cols + 1):
        col = get_column_letter(i)
        maxlen = max((len(str(c.value)) for c in ws[col] if c.value is not None), default=8)
        ws.column_dimensions[col].width = min(max(maxlen + 2, 10), 40)


# ---------------------------------------------------------------------------
# Hoja de tablas de referencia (Grubbs + Dixon) — igual que en LA-F-210/319
# ---------------------------------------------------------------------------
def _write_tablas_sheet(wb) -> str:
    ws = wb.create_sheet("Tablas")
    ws.cell(row=1, column=1, value="Tabla de valores críticos de Grubbs (rechazo de un dato)").font = BOLD
    _hdr(ws, 2, 1, "n")
    _hdr(ws, 2, 2, "G crítico (α=0.05)")
    _hdr(ws, 2, 3, "G crítico (α=0.01)")
    r = 3
    for n, vals in sorted(GRUBBS_CRITICAL.items()):
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=vals[0.05])
        ws.cell(row=r, column=3, value=vals[0.01])
        r += 1

    ws.cell(row=1, column=6, value="Tabla de valores críticos de Dixon (Q, rechazo de un dato, 95%)").font = BOLD
    _hdr(ws, 2, 6, "n")
    _hdr(ws, 2, 7, "Q crítico (95%)")
    r = 3
    for n, val in sorted(DIXON_CRITICAL_95.items()):
        ws.cell(row=r, column=6, value=n)
        ws.cell(row=r, column=7, value=val)
        r += 1
    _autosize(ws, 7)
    return ws.title


# ---------------------------------------------------------------------------
# Bloque: curva de calibración con fórmulas SLOPE/INTERCEPT/CORREL/RSQ
# ---------------------------------------------------------------------------
def _write_calibration_block(ws, row0, compound, crit, tablas_sheet):
    ws.cell(row=row0, column=1, value="CURVA DE CALIBRACIÓN — datos crudos").font = BOLD
    row = row0 + 1
    curves = [c for c in compound.calibration_curves if any(c.responses)]
    if not curves:
        return row, None
    n_niv = len(curves[0].levels_nominal)
    _hdr(ws, row, 1, "Nivel nominal")
    for i in range(len(curves)):
        _hdr(ws, row, 2 + i, f"Curva {i+1}")
    header_row = row
    row += 1
    level_col_letter = get_column_letter(1)
    first_data_row = row
    for j in range(n_niv):
        ws.cell(row=row, column=1, value=curves[0].levels_nominal[j])
        for i, curve in enumerate(curves):
            ws.cell(row=row, column=2 + i, value=curve.responses[j])
        row += 1
    last_data_row = row - 1

    n_curves = len(curves)
    y_range = f"B{first_data_row}:{get_column_letter(1+n_curves)}{last_data_row}"
    # bloque X paralelo (repite el nivel nominal en cada columna de curva)
    row += 1
    ws.cell(row=row, column=1, value="(bloque auxiliar X, igual forma que Y, para SLOPE/INTERCEPT/CORREL)").font = Font(italic=True, size=8)
    row += 1
    x_first_row = row
    for j in range(n_niv):
        for i in range(n_curves):
            _formula(ws, row, 2 + i, f"=$A{first_data_row + j}")
        row += 1
    x_last_row = row - 1
    x_range = f"B{x_first_row}:{get_column_letter(1+n_curves)}{x_last_row}"

    row += 1
    ws.cell(row=row, column=1, value="Homogeneidad de varianzas (ISO 8466-1 §4.1.2)").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="SD nivel más bajo")
    r_sbajo = row
    _formula(ws, row, 2, f"=_xlfn.STDEV.S(B{first_data_row}:{get_column_letter(1+n_curves)}{first_data_row})")
    row += 1
    ws.cell(row=row, column=1, value="SD nivel más alto")
    r_salto = row
    _formula(ws, row, 2, f"=_xlfn.STDEV.S(B{last_data_row}:{get_column_letter(1+n_curves)}{last_data_row})")
    row += 1
    ws.cell(row=row, column=1, value="Fcalc = (max/min)²")
    r_fcalc = row
    _formula(ws, row, 2, f"=(MAX(B{r_sbajo},B{r_salto})/MIN(B{r_sbajo},B{r_salto}))^2")
    row += 1
    ws.cell(row=row, column=1, value="Fcrit (f1=f2=n_curvas-1; 99%)")
    r_fcrit = row
    _formula(ws, row, 2, f"=_xlfn.F.INV(0.99,{n_curves-1},{n_curves-1})")
    row += 1
    ws.cell(row=row, column=1, value="Decisión")
    r_decision = row
    _formula(ws, row, 2, f'=IF(B{r_fcalc}>B{r_fcrit},"NO homogénea: usar PONDERADA","Homogénea: Simple es válida")')
    row += 2

    ws.cell(row=row, column=1, value="Regresión (Simple, no ponderada)").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="Pendiente")
    r_slope = row
    _formula(ws, row, 2, f"=SLOPE({y_range},{x_range})")
    row += 1
    ws.cell(row=row, column=1, value="Intercepto")
    r_intercept = row
    _formula(ws, row, 2, f"=INTERCEPT({y_range},{x_range})")
    row += 1
    ws.cell(row=row, column=1, value="r")
    r_r = row
    _formula(ws, row, 2, f"=CORREL({x_range},{y_range})")
    row += 1
    ws.cell(row=row, column=1, value="r²")
    _formula(ws, row, 2, f"=RSQ({y_range},{x_range})")
    row += 1
    ws.cell(row=row, column=1, value=f"Cumple r≥{crit.r_min}")
    _formula(ws, row, 2, f'=IF(B{r_r}>={crit.r_min},"CUMPLE","NO CUMPLE")')
    row += 2

    # Análisis completo (incluye Mandel y, si aplica, regresión ponderada) — calculado en Python
    try:
        all_lv = [c.levels_nominal for c in curves]
        all_rs = [c.responses for c in curves]
        dias_c = [c.dia or (i + 1) for i, c in enumerate(curves)]
        analistas_c = [c.analista for c in curves]
        cal = S.analyze_calibration(all_lv, all_rs, crit.r_min, dias_c, analistas_c)
        ws.cell(row=row, column=1, value="Prueba de linealidad de Mandel (ISO 8466-1 §4.1.3) y modelo final").font = BOLD
        ws.cell(row=row, column=6,
                value="⚠️ Calculado por la aplicación (ajuste cuadrático / regresión ponderada no son "
                      "fórmulas nativas simples de Excel) — ver metodología en core/stats.py").font = Font(italic=True, size=8, color="806000")
        row += 1
        ws.cell(row=row, column=1, value="PG (Mandel)")
        ws.cell(row=row, column=2, value=round(cal.linealidad_mandel.pg, 5))
        ws.cell(row=row, column=3, value="Fcrit")
        ws.cell(row=row, column=4, value=round(cal.linealidad_mandel.f_crit, 5))
        cell = ws.cell(row=row, column=5, value=cal.linealidad_mandel.decision)
        cell.fill = GREEN if cal.linealidad_mandel.es_lineal else RED
        row += 1
        ws.cell(row=row, column=1, value="Modelo final aplicado")
        cell = ws.cell(row=row, column=2, value=cal.modelo_usado)
        cell.fill = GREEN if cal.modelo_usado == "Simple" else YELLOW
        if cal.modelo_usado == "Ponderada":
            ws.cell(row=row, column=3, value="Pendiente pond.")
            ws.cell(row=row, column=4, value=round(cal.slope, 6))
            ws.cell(row=row, column=5, value="Intercepto pond.")
            ws.cell(row=row, column=6, value=round(cal.intercept, 6))
        row += 2

        ws.cell(row=row, column=1, value="Residuales y % error (según modelo final aplicado)").font = BOLD
        if cal.modelo_usado != "Simple":
            ws.cell(row=row, column=6,
                    value="⚠️ Modelo Ponderado: columnas calculadas por la app (no fórmula nativa)"
                    ).font = Font(italic=True, size=8, color="806000")
        row += 1
        _hdr(ws, row, 1, "Día"); _hdr(ws, row, 2, "Analista"); _hdr(ws, row, 3, "Nivel")
        _hdr(ws, row, 4, "Respuesta"); _hdr(ws, row, 5, "Residual")
        _hdr(ws, row, 6, "Conc. recalculada"); _hdr(ws, row, 7, "% Error")
        row += 1
        err_max = crit.error_rel_max_percent
        for dia, analista, x, y, res, xr, err in zip(
            cal.dias, cal.analistas, cal.puntos_x, cal.puntos_y, cal.residuales, cal.x_recalculada, cal.error_percent
        ):
            ws.cell(row=row, column=1, value=dia)
            ws.cell(row=row, column=2, value=analista)
            ws.cell(row=row, column=3, value=x)
            ws.cell(row=row, column=4, value=y)
            if cal.modelo_usado == "Simple":
                # fórmulas vivas: referencian la pendiente/intercepto ya calculados arriba (Simple)
                _formula(ws, row, 5, f"=D{row}-($B${r_slope}*C{row}+$B${r_intercept})")
                _formula(ws, row, 6, f"=(D{row}-$B${r_intercept})/$B${r_slope}")
                _formula(ws, row, 7, f"=((F{row}-C{row})/C{row})*100")
            else:
                ws.cell(row=row, column=5, value=round(res, 5))
                ws.cell(row=row, column=6, value=round(xr, 5))
                ws.cell(row=row, column=7, value=round(err, 2))
            c = ws.cell(row=row, column=7)
            if abs(err) > err_max:
                c.fill = RED
            row += 1
    except Exception as e:
        ws.cell(row=row, column=1, value=f"(No se pudo completar el análisis Mandel/ponderado: {e})")
        row += 1

    return row + 1, {"slope_cell": f"B{r_slope}", "intercept_cell": f"B{r_intercept}"}


# ---------------------------------------------------------------------------
# Bloque: un nivel (LC, LS, muestra) — datos crudos + fórmulas completas
# ---------------------------------------------------------------------------
def _write_level_block(ws, row0, lvl, design, crit, label_extra=""):
    d = design.n_grupos
    r = design.replicas_por_grupo
    grupo_labels = [f"Día {g['dia']} — {g['analista']}" for g in design.grupos]

    ws.cell(row=row0, column=1, value=f"{lvl.label}{label_extra} — datos crudos (nominal = {lvl.nominal})").font = BOLD
    row = row0 + 1
    _hdr(ws, row, 1, "Grupo")
    for j in range(r):
        _hdr(ws, row, 2 + j, f"Rép. {j+1}")
    _hdr(ws, row, 2 + r, "Suma fila")
    row += 1
    first_row = row
    for i, glabel in enumerate(grupo_labels):
        ws.cell(row=row, column=1, value=glabel)
        vals = lvl.values[i] if i < len(lvl.values) else [0.0] * r
        for j in range(r):
            ws.cell(row=row, column=2 + j, value=vals[j] if j < len(vals) else 0.0)
        rl = get_column_letter(2)
        rl2 = get_column_letter(1 + r)
        _formula(ws, row, 2 + r, f"=SUM({rl}{row}:{rl2}{row})")
        row += 1
    last_row = row - 1
    data_range = f"B{first_row}:{get_column_letter(1+r)}{last_row}"
    sum_col_range = f"{get_column_letter(2+r)}{first_row}:{get_column_letter(2+r)}{last_row}"

    row += 1
    ws.cell(row=row, column=1, value="RESULTADOS ESTADÍSTICOS").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="n")
    r_n = row
    _formula(ws, row, 2, f"=COUNT({data_range})")
    row += 1
    ws.cell(row=row, column=1, value="Promedio")
    r_mean = row
    _formula(ws, row, 2, f"=AVERAGE({data_range})")
    row += 1
    ws.cell(row=row, column=1, value="Desviación estándar")
    r_sd = row
    _formula(ws, row, 2, f"=_xlfn.STDEV.S({data_range})")
    row += 1
    ws.cell(row=row, column=1, value="Coeficiente de variación, %")
    r_cv = row
    _formula(ws, row, 2, f"=(B{r_sd}/B{r_mean})*100")
    c = ws.cell(row=row, column=3, value=f'=IF(B{r_cv}<={crit.cv_max_percent},"CUMPLE","NO CUMPLE")')
    c.font = FORMULA_FONT
    row += 1

    r_error = r_recov = r_tcalc = r_ttab = None
    if lvl.nominal:
        ws.cell(row=row, column=1, value="Valor nominal")
        r_nom = row
        ws.cell(row=row, column=2, value=lvl.nominal)
        row += 1
        ws.cell(row=row, column=1, value="Error absoluto")
        r_errabs = row
        _formula(ws, row, 2, f"=B{r_mean}-B{r_nom}")
        row += 1
        ws.cell(row=row, column=1, value="Error relativo, %")
        r_error = row
        _formula(ws, row, 2, f"=(B{r_errabs}/B{r_nom})*100")
        row += 1
        ws.cell(row=row, column=1, value="Recuperación, %")
        r_recov = row
        _formula(ws, row, 2, f"=(B{r_mean}/B{r_nom})*100")
        rec_min, rec_max = (crit.recovery_lc_min, crit.recovery_lc_max) if "LC" in lvl.label else (crit.recovery_min, crit.recovery_max)
        c = ws.cell(row=row, column=3, value=f'=IF(AND(B{r_recov}>={rec_min},B{r_recov}<={rec_max}),"CUMPLE","NO CUMPLE")')
        c.font = FORMULA_FONT
        row += 1
        ws.cell(row=row, column=1, value="t calculado")
        r_tcalc = row
        _formula(ws, row, 2, f"=ABS(B{r_errabs})/(B{r_sd}/SQRT(B{r_n}))")
        row += 1
        ws.cell(row=row, column=1, value="t tabulado (dos colas)")
        r_ttab = row
        _formula(ws, row, 2, f"=_xlfn.T.INV.2T({crit.t_alpha},B{r_n}-1)")
        c = ws.cell(row=row, column=3, value=f'=IF(B{r_tcalc}<=B{r_ttab},"Sin sesgo significativo","SESGO SIGNIFICATIVO")')
        c.font = FORMULA_FONT
        row += 1

    # ANOVA anidado (Sr / SI) — Eurachem 6.6.4 / ISO 5725-3
    row += 1
    ws.cell(row=row, column=1, value="ANOVA anidado (repetibilidad Sr / precisión intermedia SI)").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="N total")
    r_N = row
    _formula(ws, row, 2, f"=B{r_n}")
    row += 1
    ws.cell(row=row, column=1, value="CT (término de corrección)")
    r_CT = row
    _formula(ws, row, 2, f"=(SUM({data_range}))^2/B{r_N}")
    row += 1
    ws.cell(row=row, column=1, value="SS total")
    r_SStot = row
    _formula(ws, row, 2, f"=SUMSQ({data_range})-B{r_CT}")
    row += 1
    ws.cell(row=row, column=1, value="SS entre grupos")
    r_SSentre = row
    _formula(ws, row, 2, f"=SUMSQ({sum_col_range})/{r}-B{r_CT}")
    row += 1
    ws.cell(row=row, column=1, value="SS dentro de grupos")
    r_SSdentro = row
    _formula(ws, row, 2, f"=B{r_SStot}-B{r_SSentre}")
    row += 1
    ws.cell(row=row, column=1, value="df entre / dentro")
    r_dfentre = row
    ws.cell(row=row, column=2, value=d - 1)
    ws.cell(row=row, column=3, value=d * (r - 1))
    r_dfdentro_col = 3
    row += 1
    ws.cell(row=row, column=1, value="MS entre / MS dentro")
    r_MSentre = row
    _formula(ws, row, 2, f"=B{r_SSentre}/B{r_dfentre}")
    _formula(ws, row, 3, f"=B{r_SSdentro}/C{r_dfentre}")
    r_MSdentro_col = 3
    r_MSdentro_row = row
    row += 1
    ws.cell(row=row, column=1, value="Sr (repetibilidad)")
    r_sr = row
    _formula(ws, row, 2, f"=SQRT(MAX(C{r_MSdentro_row},0))")
    row += 1
    ws.cell(row=row, column=1, value="SI (precisión intermedia)")
    r_si = row
    _formula(ws, row, 2, f"=SQRT(MAX((B{r_MSentre}-C{r_MSdentro_row})/{r},0)+B{r_sr}^2)")
    row += 1
    ws.cell(row=row, column=1, value="% Sr / % SI")
    _formula(ws, row, 2, f"=(B{r_sr}/B{r_mean})*100")
    _formula(ws, row, 3, f"=(B{r_si}/B{r_mean})*100")
    row += 1

    # Dato atípico: Dixon (n<=7) o Grubbs (n>=8), automático
    row += 1
    ws.cell(row=row, column=1, value="Datos atípicos (Dixon si n≤7, Grubbs si n≥8 — automático)").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="Prueba usada")
    r_test = row
    _formula(ws, row, 2, f'=IF(B{r_n}<4,"N/A",IF(B{r_n}<=7,"Dixon","Grubbs"))')
    row += 1
    ws.cell(row=row, column=1, value="Valor crítico")
    r_crit = row
    grubbs_col = 3 if crit.grubbs_alpha == 0.01 else 2
    _formula(ws, row, 2,
             f'=IF(B{r_test}="Dixon",VLOOKUP(B{r_n},Tablas!$F$3:$G$19,2,FALSE),'
             f'IF(B{r_test}="Grubbs",VLOOKUP(B{r_n},Tablas!$A$3:$C$30,{grubbs_col},FALSE),""))')
    row += 1
    ws.cell(row=row, column=1, value="Estadístico bajo / alto")
    r_stat = row
    _formula(ws, row, 2,
             f'=IF(B{r_test}="Dixon",(SMALL({data_range},2)-MIN({data_range}))/(MAX({data_range})-MIN({data_range})),'
             f'IF(B{r_test}="Grubbs",(B{r_mean}-MIN({data_range}))/B{r_sd},""))')
    _formula(ws, row, 3,
             f'=IF(B{r_test}="Dixon",(MAX({data_range})-LARGE({data_range},2))/(MAX({data_range})-MIN({data_range})),'
             f'IF(B{r_test}="Grubbs",(MAX({data_range})-B{r_mean})/B{r_sd},""))')
    row += 1
    ws.cell(row=row, column=1, value="Decisión")
    c = ws.cell(row=row, column=2,
                value=f'=IF(OR(B{r_stat}>B{r_crit},C{r_stat}>B{r_crit}),"POSIBLE ATÍPICO","Sin atípicos")')
    c.font = FORMULA_FONT
    row += 1

    refs = {"mean": f"B{r_mean}", "sd": f"B{r_sd}", "cv": f"B{r_cv}", "n": f"B{r_n}",
            "si": f"B{r_si}", "sr": f"B{r_sr}"}
    return row + 1, refs


def build_verification_excel(project: VerificationProject, out_path: str = "data/Verificacion.xlsx") -> str:
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    crit = project.criteria
    tablas_sheet = _write_tablas_sheet(wb)

    for compound in project.compounds:
        ws = wb.create_sheet(compound.nombre[:31])
        row = 1
        ws.cell(row=row, column=1, value=f"VERIFICACIÓN — {compound.nombre}").font = Font(bold=True, size=14)
        row += 1
        ws.cell(row=row, column=1,
                value=f"Ensayo: {project.method.ensayo}  |  Técnica: {project.method.tecnica}  |  "
                      f"Matriz: {project.method.matriz}  |  Celdas en azul = fórmula viva")
        row += 2

        row, curve_refs = _write_calibration_block(ws, row, compound, crit, tablas_sheet)
        row += 1

        level_refs = {}
        for lvl in compound.replicate_levels:
            if not any(v for day in lvl.values for v in day):
                continue
            row, refs = _write_level_block(ws, row, lvl, project.design, crit)
            level_refs[lvl.label] = refs
            row += 1

        # --- Incertidumbre (se mantiene como valores calculados: combina
        # múltiples fuentes -- prep. estándar, certificados -- que no viven
        # como celdas en esta hoja; se referencia el Sr/SI ya calculado arriba
        # cuando corresponde) ---
        if compound.standard_preparation and compound.standard_preparation.dilution_levels:
            ws.cell(row=row, column=1, value="INCERTIDUMBRE DE MEDICIÓN (QUAM:2012 CG4)").font = BOLD
            ws.cell(row=row, column=6, value="⚠️ Valores calculados por la app (combina fuentes externas: "
                                              "certificados MRC, material volumétrico)").font = Font(italic=True, size=8, color="806000")
            row += 1
            _hdr(ws, row, 1, "Nivel"); _hdr(ws, row, 2, "Concentración"); _hdr(ws, row, 3, "u cal")
            _hdr(ws, row, 4, "u prep"); _hdr(ws, row, 5, "u vol.muestra"); _hdr(ws, row, 6, "u repetibilidad")
            _hdr(ws, row, 7, "uc"); _hdr(ws, row, 8, "U expandida"); _hdr(ws, row, 9, "U relativa %")
            _hdr(ws, row, 10, "Resultado")
            row += 1
            u_vm = U.volumetric_step_uncertainty(compound.volumen_muestra) if compound.volumen_muestra else None
            u_vm_rel = u_vm.u_relativa if u_vm else 0.0
            all_lv_e = [c.levels_nominal for c in compound.calibration_curves if any(c.responses)]
            all_rs_e = [c.responses for c in compound.calibration_curves if any(c.responses)]
            cal_e = None
            if len(all_lv_e) >= 2 and all(len(l) >= 4 for l in all_lv_e):
                try:
                    cal_e = S.analyze_calibration(all_lv_e, all_rs_e, crit.r_min)
                except Exception:
                    cal_e = None
            for label, nominal in [("LC", compound.lc_nominal), ("LS", compound.ls_nominal)]:
                lvl = next((r for r in compound.replicate_levels if r.label == label), None)
                if lvl is None:
                    continue
                flat = [v for day in lvl.values for v in day]
                if not any(flat):
                    continue
                try:
                    prep_u = U.standard_preparation_uncertainty_at(compound.standard_preparation, nominal)
                except Exception:
                    continue
                try:
                    prec = S.repeatability_intermediate_precision(lvl.values)
                    u_rep_rel = prec.si / (sum(flat) / len(flat))
                except Exception:
                    u_rep_rel = 0.0
                try:
                    if cal_e is None:
                        raise ValueError("Sin modelo de calibración")
                    res_x = [rr / cal_e.slope for rr in cal_e.residuales]
                    u_cal_abs = U.calibration_response_uncertainty(cal_e.puntos_x, res_x, nominal)
                    u_cal_rel = u_cal_abs / nominal if nominal else 0.0
                except Exception:
                    u_cal_rel = 0.0
                budget = U.combine_budget(label, nominal, u_cal_rel, prep_u.u_relativa_combinada,
                                           u_vm_rel, u_rep_rel, k=crit.k_coverage)
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=budget.concentracion)
                ws.cell(row=row, column=3, value=round(budget.u_calibracion_relativa, 5))
                ws.cell(row=row, column=4, value=round(budget.u_preparacion_relativa, 5))
                ws.cell(row=row, column=5, value=round(budget.u_volumen_muestra_relativa, 5))
                ws.cell(row=row, column=6, value=round(budget.u_repetibilidad_relativa, 5))
                ws.cell(row=row, column=7, value=round(budget.uc_relativa, 5))
                ws.cell(row=row, column=8, value=round(budget.u_expandida, 5))
                ws.cell(row=row, column=9, value=round(budget.u_relativa_expandida_percent, 2))
                ws.cell(row=row, column=10, value=f"({budget.concentracion} ± {round(budget.u_expandida,5)}) {project.method.unidad}")
                row += 1

        _autosize(ws, 12)

    wb.save(out_path)
    return out_path
